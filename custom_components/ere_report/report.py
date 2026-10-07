"""Build the quarterly report and write it as xlsx and csv.

Pure Python plus openpyxl: no Home Assistant imports. The report is written
in Dutch by default (the ERE scheme is Dutch) or in English.
"""

from __future__ import annotations

import csv
from dataclasses import dataclass, field, replace
from datetime import UTC, datetime, tzinfo
from pathlib import Path

from .history import (
    HOUR,
    SUSPECT_KWH_PER_HOUR,
    HourRow,
    quarter_bounds,
    reconstruct_sessions,
)
from .report_text import DEFAULT_LANGUAGE, texts
from .session_tracker import (
    FLAG_ONGOING,
    FLAG_SPLIT,
    SOURCE_RECONSTRUCTED,
    SOURCE_UNOBSERVED,
    Session,
)

DT_FORMAT = "DD-MM-YYYY HH:MM"


@dataclass
class ReportMeta:
    """Who and what the report is about."""

    charger_name: str
    energy_entity: str
    holder_name: str = ""
    address: str = ""
    postcode_city: str = ""
    ean: str = ""
    charger_brand: str = ""
    charger_model: str = ""
    charger_serial: str = ""


@dataclass
class ReportData:
    """Everything that ends up in the report files."""

    meta: ReportMeta
    year: int
    quarter: int
    start: datetime
    end: datetime
    generated: datetime
    complete: bool
    meter_begin: float | None
    meter_end: float | None
    meter_total: float
    sessions: list[Session]
    month_meter_kwh: dict[int, float]
    language: str = DEFAULT_LANGUAGE
    notes: list[str] = field(default_factory=list)

    @property
    def sessions_total(self) -> float:
        return round(sum(s.kwh for s in self.sessions), 2)

    @property
    def unallocated(self) -> float:
        return round(self.meter_total - self.sessions_total, 2)

    @property
    def months(self) -> list[int]:
        first = 3 * (self.quarter - 1) + 1
        return [first, first + 1, first + 2]


def _clip(
    session: Session,
    start: datetime,
    end: datetime,
    meter_begin: float | None,
    meter_end: float | None,
) -> Session | None:
    """Limit a session to the report period, using the boundary meter readings."""
    if session.end <= start or session.start >= end:
        return None
    if session.start < start:
        if meter_begin is None or session.meter_end is None:
            return None
        kwh = max(0.0, min(session.kwh, session.meter_end - meter_begin))
        session = replace(
            session,
            start=start,
            meter_start=meter_begin,
            kwh=round(kwh, 3),
            flags=[*session.flags, FLAG_SPLIT],
        )
    if session.end > end:
        if meter_end is None or session.meter_start is None:
            return None
        kwh = max(0.0, min(session.kwh, meter_end - session.meter_start))
        flags = [f for f in session.flags if f != FLAG_ONGOING]
        session = replace(
            session,
            end=end,
            meter_end=meter_end,
            kwh=round(kwh, 3),
            flags=[*dict.fromkeys([*flags, FLAG_SPLIT])],
        )
    return session if session.kwh > 0 else None


def build_report(
    meta: ReportMeta,
    year: int,
    quarter: int,
    tz: tzinfo,
    now: datetime,
    rows: list[HourRow],
    recorded_sessions: list[Session],
    tracking_since: datetime | None,
    min_kwh: float,
    language: str = DEFAULT_LANGUAGE,
) -> ReportData:
    """Combine statistics and recorded sessions into a report.

    ``rows`` are hourly statistics covering the quarter plus some hours before
    it, so the meter reading at the start of the quarter can be determined.
    """
    start, end = quarter_bounds(year, quarter, tz)
    rows = sorted(rows, key=lambda r: r.start)
    before = [r for r in rows if r.start < start]
    period = [r for r in rows if start <= r.start < end]
    complete = now >= end

    meter_total = round(sum(r.change for r in period), 2)
    meter_begin = None
    if before and before[-1].state is not None:
        meter_begin = before[-1].state
    elif period and period[0].state is not None:
        meter_begin = period[0].state - period[0].change
    meter_end = period[-1].state if period else None
    if meter_begin is not None:
        meter_begin = round(meter_begin, 3)
    if meter_end is not None:
        meter_end = round(meter_end, 3)

    reconstruct_rows = [
        r for r in period if tracking_since is None or r.start + HOUR <= tracking_since
    ]
    sessions = reconstruct_sessions(reconstruct_rows, min_kwh)
    for recorded in recorded_sessions:
        if clipped := _clip(recorded, start, end, meter_begin, meter_end):
            sessions.append(clipped)
    sessions.sort(key=lambda s: s.start)

    month_meter_kwh: dict[int, float] = {}
    for row in period:
        month = row.start.astimezone(tz).month
        month_meter_kwh[month] = month_meter_kwh.get(month, 0.0) + row.change

    data = ReportData(
        meta=meta,
        year=year,
        quarter=quarter,
        start=start,
        end=end,
        generated=now,
        complete=complete,
        meter_begin=meter_begin,
        meter_end=meter_end,
        meter_total=meter_total,
        sessions=sessions,
        month_meter_kwh=month_meter_kwh,
        language=language,
    )
    data.notes = _notes(data, period, now, tz)
    return data


def _notes(
    data: ReportData, period: list[HourRow], now: datetime, tz: tzinfo
) -> list[str]:
    t = texts(data.language)
    notes = []
    if not data.complete:
        notes.append(t["note_incomplete"])
    if not period:
        notes.append(t["note_no_data"])
        return notes

    # Subtract in UTC: local wall-clock arithmetic ignores DST changes.
    period_end = min(data.end, now).astimezone(UTC)
    expected = int((period_end - data.start.astimezone(UTC)) / HOUR)
    if (missing := expected - len(period)) > 0:
        notes.append(t["note_missing"].format(missing=missing, expected=expected))
    if negative := [r for r in period if r.change < 0]:
        notes.append(t["note_negative"].format(count=len(negative)))
    for row in period:
        if row.change > SUSPECT_KWH_PER_HOUR:
            when = row.start.astimezone(tz).strftime("%d-%m-%Y %H:%M")
            notes.append(t["note_spike"].format(kwh=row.change, when=when))
    if data.meter_begin is not None and data.meter_end is not None:
        by_readings = round(data.meter_end - data.meter_begin, 2)
        if abs(by_readings - data.meter_total) > 0.02:
            notes.append(
                t["note_mismatch"].format(
                    by_readings=by_readings, total=data.meter_total
                )
            )
    counts = {
        source: sum(1 for s in data.sessions if s.source == source)
        for source in (SOURCE_RECONSTRUCTED, SOURCE_UNOBSERVED)
    }
    if counts[SOURCE_RECONSTRUCTED]:
        notes.append(t["note_reconstructed"].format(count=counts[SOURCE_RECONSTRUCTED]))
    if counts[SOURCE_UNOBSERVED]:
        notes.append(t["note_unobserved"].format(count=counts[SOURCE_UNOBSERVED]))
    return notes


def _period_label(data: ReportData, tz: tzinfo) -> str:
    t = texts(data.language)
    months = t["months"]
    last_day = (data.end - HOUR).astimezone(tz)
    first_day = data.start.astimezone(tz)
    return (
        f"{first_day.day} {months[first_day.month - 1]} {first_day.year} "
        f"{t['period_joiner']} "
        f"{last_day.day} {months[last_day.month - 1]} {last_day.year} ({tz})"
    )


def _source_label(data: ReportData, session: Session) -> str:
    return texts(data.language)["sources"].get(session.source, session.source)


def _remarks(data: ReportData, session: Session) -> str:
    labels = texts(data.language)["flags"]
    return "; ".join(labels.get(flag, flag) for flag in session.flags)


def write_xlsx(data: ReportData, path: Path, tz: tzinfo, version: str) -> None:
    """Write the report as an Excel workbook with plain values (no formulas)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    title_font = Font(bold=True, size=14)
    bold = Font(bold=True)
    header_fill = PatternFill("solid", start_color="DDEBF7")

    def header(ws, row: int, labels: list[str]) -> None:
        for col, label in enumerate(labels, start=1):
            cell = ws.cell(row=row, column=col, value=label)
            cell.font = bold
            cell.fill = header_fill

    t = texts(data.language)
    wb = Workbook()
    ws = wb.active
    ws.title = t["sheet_summary"]
    meta = data.meta
    ws.append([t["title"].format(quarter=data.quarter, year=data.year)])
    ws["A1"].font = title_font
    ws.append([_period_label(data, tz)])
    ws.append([])

    sections: list[tuple[str, list[tuple[str, object]]]] = [
        (
            t["section_applicant"],
            [
                (t["name"], meta.holder_name),
                (t["address"], meta.address),
                (t["postcode_city"], meta.postcode_city),
                (t["ean"], meta.ean),
            ],
        ),
        (
            t["section_charger"],
            [
                (t["name"], meta.charger_name),
                (t["brand"], meta.charger_brand),
                (t["model"], meta.charger_model),
                (t["serial"], meta.charger_serial),
                (t["source_entity"], meta.energy_entity),
            ],
        ),
        (
            t["section_readings"],
            [
                (t["meter_begin"], data.meter_begin),
                (t["meter_end"], data.meter_end),
                (t["delivered"], data.meter_total),
                (t["session_count"], len(data.sessions)),
                (t["sessions_total"], data.sessions_total),
                (t["unallocated"], data.unallocated),
            ],
        ),
    ]
    for title, items in sections:
        ws.append([title])
        ws.cell(row=ws.max_row, column=1).font = bold
        ws.cell(row=ws.max_row, column=1).fill = header_fill
        ws.cell(row=ws.max_row, column=2).fill = header_fill
        for label, value in items:
            ws.append([label, value])
            cell = ws.cell(row=ws.max_row, column=2)
            cell.alignment = Alignment(horizontal="left")
            if label == t["delivered"]:
                ws.cell(row=ws.max_row, column=1).font = bold
                cell.font = bold
            if isinstance(value, float):
                cell.number_format = "0.00"
        ws.append([])
    if data.notes:
        ws.append([t["section_notes"]])
        ws.cell(row=ws.max_row, column=1).font = bold
        for note in data.notes:
            ws.append([note])
        ws.append([])
    generated = data.generated.astimezone(tz).strftime("%d-%m-%Y %H:%M")
    ws.append([t["footer"].format(generated=generated, version=version)])
    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 36

    ws = wb.create_sheet(t["sheet_sessions"])
    header(ws, 1, t["session_headers"])
    for number, session in enumerate(data.sessions, start=1):
        local_start = session.start.astimezone(tz)
        local_end = session.end.astimezone(tz)
        ws.append(
            [
                number,
                t["days"][local_start.weekday()],
                local_start.replace(tzinfo=None),
                local_end.replace(tzinfo=None),
                round((session.end - session.start) / HOUR, 2),
                session.meter_start,
                session.meter_end,
                round(session.kwh, 2),
                _source_label(data, session),
                _remarks(data, session),
            ]
        )
        row = ws.max_row
        ws.cell(row=row, column=3).number_format = DT_FORMAT
        ws.cell(row=row, column=4).number_format = DT_FORMAT
        for col in (5, 8):
            ws.cell(row=row, column=col).number_format = "0.00"
        for col in (6, 7):
            ws.cell(row=row, column=col).number_format = "0.000"
    ws.append([])
    ws.append([t["total"], None, None, None, None, None, None, data.sessions_total])
    ws.cell(row=ws.max_row, column=1).font = bold
    ws.cell(row=ws.max_row, column=8).font = bold
    ws.cell(row=ws.max_row, column=8).number_format = "0.00"
    for column, width in zip(
        "ABCDEFGHIJ", (6, 6, 18, 18, 9, 22, 22, 10, 30, 40), strict=True
    ):
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A2"

    ws = wb.create_sheet(t["sheet_months"])
    header(ws, 1, t["month_headers"])
    for month in data.months:
        in_month = [s for s in data.sessions if s.start.astimezone(tz).month == month]
        ws.append(
            [
                t["months"][month - 1],
                len(in_month),
                round(sum(s.kwh for s in in_month), 2),
                round(data.month_meter_kwh.get(month, 0.0), 2),
            ]
        )
    ws.append([])
    ws.append([t["total"], len(data.sessions), data.sessions_total, data.meter_total])
    for col in range(1, 5):
        ws.cell(row=ws.max_row, column=col).font = bold
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=4):
        for cell in row:
            cell.number_format = "0.00"
    for column, width in zip("ABCD", (14, 16, 14, 14), strict=True):
        ws.column_dimensions[column].width = width

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_csv(data: ReportData, path: Path, tz: tzinfo) -> None:
    """Write one line per session, with ISO 8601 timestamps."""
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(texts(data.language)["csv_columns"])
        for session in data.sessions:
            writer.writerow(
                [
                    data.meta.charger_name,
                    data.meta.charger_serial,
                    data.meta.ean,
                    session.start.astimezone(tz).isoformat(timespec="seconds"),
                    session.end.astimezone(tz).isoformat(timespec="seconds"),
                    "" if session.meter_start is None else f"{session.meter_start:.3f}",
                    "" if session.meter_end is None else f"{session.meter_end:.3f}",
                    f"{session.kwh:.2f}",
                    _source_label(data, session),
                    _remarks(data, session),
                ]
            )
