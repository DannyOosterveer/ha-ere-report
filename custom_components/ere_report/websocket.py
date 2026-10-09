"""WebSocket API for the reports panel."""

from __future__ import annotations

from datetime import UTC, datetime
from pathlib import Path
import re
from typing import Any

from homeassistant.components import websocket_api
from homeassistant.config_entries import ConfigEntryState
from homeassistant.core import HomeAssistant, callback
from homeassistant.util import slugify
import voluptuous as vol

from .const import DOMAIN, REPORT_DIR
from .report_text import TEXTS

REPORT_STEM = re.compile(
    r"ere_(?P<slug>[a-z0-9_]+)_(?P<year>\d{4})_q(?P<quarter>[1-4])"
)
REPORT_FILENAME = re.compile(REPORT_STEM.pattern + r"\.(?P<ext>xlsx|csv)")


@callback
def async_register(hass: HomeAssistant) -> None:
    websocket_api.async_register_command(hass, ws_reports)
    websocket_api.async_register_command(hass, ws_delete)


def _scan(folder: Path) -> dict[str, dict[str, Any]]:
    """Group the report files in the folder by stem."""
    found: dict[str, dict[str, Any]] = {}
    if not folder.is_dir():
        return found
    for path in folder.iterdir():
        if not (match := REPORT_FILENAME.fullmatch(path.name)):
            continue
        stem = path.name.rsplit(".", 1)[0]
        modified = datetime.fromtimestamp(path.stat().st_mtime, UTC)
        item = found.setdefault(
            stem,
            {
                "stem": stem,
                "slug": match["slug"],
                "year": int(match["year"]),
                "quarter": int(match["quarter"]),
                "modified": modified,
                "files": [],
            },
        )
        item["files"].append(match["ext"])
        item["modified"] = max(item["modified"], modified)
    return found


def _read_summary(path: Path) -> dict[str, Any]:
    """Read the delivered kWh and session count from a report's summary sheet.

    Used for reports created before the integration kept a record of them.
    """
    from openpyxl import load_workbook

    labels = {
        **{t["delivered"]: "total_kwh" for t in TEXTS.values()},
        **{t["session_count"]: "sessions" for t in TEXTS.values()},
    }
    found: dict[str, Any] = {}
    try:
        workbook = load_workbook(path, read_only=True)
    except Exception:  # noqa: BLE001 - a damaged file should not break the list
        return found
    try:
        for row in workbook.worksheets[0].iter_rows(max_col=2, values_only=True):
            if row and row[0] in labels and isinstance(row[1], (int, float)):
                found[labels[row[0]]] = row[1]
    finally:
        workbook.close()
    return found


@websocket_api.websocket_command({vol.Required("type"): f"{DOMAIN}/reports"})
@websocket_api.require_admin
@websocket_api.async_response
async def ws_reports(
    hass: HomeAssistant,
    connection: websocket_api.ActiveConnection,
    msg: dict[str, Any],
) -> None:
    """List charge points and the report files on disk."""
    found = await hass.async_add_executor_job(_scan, Path(hass.config.path(REPORT_DIR)))
    folder = Path(hass.config.path(REPORT_DIR))
    chargers = []
    known: dict[str, tuple[str, str, dict[str, Any]]] = {}
    by_slug: dict[str, tuple[str, str]] = {}
    for entry in hass.config_entries.async_entries(DOMAIN):
        if entry.state is not ConfigEntryState.LOADED:
            continue
        manager = entry.runtime_data
        name = manager.charger_name
        chargers.append({"entry_id": entry.entry_id, "title": name})
        for known_name in (entry.title, name):
            by_slug[slugify(known_name)] = (entry.entry_id, name)
        for stem, report in manager.reports.items():
            known[stem] = (entry.entry_id, name, report)

    reports = []
    for stem, item in found.items():
        if stem in known:
            entry_id, title, meta = known[stem]
        else:
            entry_id, title = by_slug.get(item["slug"], (None, item["slug"]))
            meta = {}
            if "xlsx" in item["files"]:
                meta = await hass.async_add_executor_job(
                    _read_summary, folder / f"{stem}.xlsx"
                )
        reports.append(
            {
                "stem": stem,
                "entry_id": entry_id,
                "charger": title,
                "year": item["year"],
                "quarter": item["quarter"],
                "generated": meta.get("generated", item["modified"].isoformat()),
                "total_kwh": meta.get("total_kwh"),
                "sessions": meta.get("sessions"),
                "complete": meta.get("complete"),
                "files": sorted(item["files"], reverse=True),
            }
        )
    reports.sort(key=lambda r: (r["year"], r["quarter"], r["charger"]), reverse=True)
    connection.send_result(msg["id"], {"chargers": chargers, "reports": reports})


def _delete_files(folder: Path, stem: str) -> list[str]:
    removed = []
    for ext in ("xlsx", "csv"):
        path = folder / f"{stem}.{ext}"
        if path.is_file():
            path.unlink()
            removed.append(ext)
    return removed


@websocket_api.websocket_command(
    {vol.Required("type"): f"{DOMAIN}/delete", vol.Required("stem"): str}
)
@websocket_api.require_admin
@websocket_api.async_response
async def ws_delete(
    hass: HomeAssistant,
    connection: websocket_api.ActiveConnection,
    msg: dict[str, Any],
) -> None:
    """Delete a report's xlsx and csv file."""
    if not (match := REPORT_STEM.fullmatch(msg["stem"])):
        connection.send_error(
            msg["id"], websocket_api.ERR_INVALID_FORMAT, "Invalid report name"
        )
        return
    removed = await hass.async_add_executor_job(
        _delete_files, Path(hass.config.path(REPORT_DIR)), msg["stem"]
    )
    if not removed:
        connection.send_error(msg["id"], websocket_api.ERR_NOT_FOUND, "No such report")
        return
    for entry in hass.config_entries.async_entries(DOMAIN):
        if entry.state is ConfigEntryState.LOADED and entry.runtime_data.owns_stem(
            msg["stem"]
        ):
            entry.runtime_data.forget_report(
                msg["stem"], int(match["year"]), int(match["quarter"])
            )
    connection.send_result(msg["id"], {"removed": removed})
