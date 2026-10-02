"""End-to-end tests: live tracking, statistics and report generation."""

from datetime import datetime, timedelta
from pathlib import Path

from freezegun.api import FrozenDateTimeFactory
from homeassistant.components.recorder.models import StatisticMeanType
from homeassistant.components.recorder.statistics import async_import_statistics
from homeassistant.const import STATE_OFF, STATE_ON, STATE_UNAVAILABLE
from homeassistant.core import HomeAssistant
from homeassistant.util import dt as dt_util
from openpyxl import load_workbook
from pytest_homeassistant_custom_component.common import (
    MockConfigEntry,
    async_fire_time_changed,
)
from pytest_homeassistant_custom_component.components.recorder.common import (
    async_wait_recording_done,
)
from pytest_homeassistant_custom_component.typing import ClientSessionGenerator

from custom_components.ere_report.const import (
    CONF_EAN,
    CONF_ENERGY_ENTITY,
    CONF_MID_CONFIRMED,
    DOMAIN,
    EVENT_REPORT_GENERATED,
)
from custom_components.ere_report.history import previous_quarter, quarter_bounds
from custom_components.ere_report.session_tracker import SOURCE_UNOBSERVED

ENTITY = "sensor.charger_meter"
ATTRS = {
    "unit_of_measurement": "kWh",
    "device_class": "energy",
    "state_class": "total_increasing",
}


async def setup_entry(hass: HomeAssistant) -> MockConfigEntry:
    entry = MockConfigEntry(
        domain=DOMAIN,
        title="Laadpaal",
        unique_id=ENTITY,
        data={CONF_ENERGY_ENTITY: ENTITY},
        options={CONF_EAN: "871234567890123456", CONF_MID_CONFIRMED: True},
    )
    entry.add_to_hass(hass)
    assert await hass.config_entries.async_setup(entry.entry_id)
    await hass.async_block_till_done()
    return entry


async def advance(
    hass: HomeAssistant, freezer: FrozenDateTimeFactory, minutes: float
) -> None:
    freezer.tick(timedelta(minutes=minutes))
    async_fire_time_changed(hass)
    await hass.async_block_till_done()


async def import_hourly(
    hass: HomeAssistant, start: datetime, changes: list[float], meter: float
) -> None:
    """Store hourly long-term statistics for the meter sensor."""
    stats = []
    total = 0.0
    for index, change in enumerate(changes):
        meter += change
        total += change
        stats.append(
            {"start": start + timedelta(hours=index), "state": meter, "sum": total}
        )
    metadata = {
        "has_sum": True,
        "mean_type": StatisticMeanType.NONE,
        "name": None,
        "source": "recorder",
        "statistic_id": ENTITY,
        "unit_class": "energy",
        "unit_of_measurement": "kWh",
    }
    async_import_statistics(hass, metadata, stats)
    await async_wait_recording_done(hass)


async def test_live_session_and_entities(
    recorder_mock, hass: HomeAssistant, freezer: FrozenDateTimeFactory
) -> None:
    hass.states.async_set(ENTITY, "1000.0", ATTRS)
    entry = await setup_entry(hass)
    manager = entry.runtime_data
    active = "binary_sensor.laadpaal_session_active"
    assert hass.states.get(active).state == STATE_OFF

    await advance(hass, freezer, 5)
    hass.states.async_set(ENTITY, "1001.5", ATTRS)
    await hass.async_block_till_done()
    assert hass.states.get(active).state == STATE_ON

    await advance(hass, freezer, 10)
    hass.states.async_set(ENTITY, "1004.0", ATTRS)
    await hass.async_block_till_done()
    await advance(hass, freezer, 10)
    hass.states.async_set(ENTITY, "1008.0", ATTRS)
    await hass.async_block_till_done()
    await advance(hass, freezer, 16)

    assert hass.states.get(active).state == STATE_OFF
    [session] = manager.sessions
    assert session.kwh == 8.0
    assert session.meter_start == 1000.0
    assert session.meter_end == 1008.0
    assert session.end - session.start == timedelta(minutes=20)
    assert hass.states.get("sensor.laadpaal_last_session_energy").state == "8.0"
    assert hass.states.get("sensor.laadpaal_sessions_this_quarter").state == "1"


async def test_wh_sensor_and_unavailable_gap(
    recorder_mock, hass: HomeAssistant, freezer: FrozenDateTimeFactory
) -> None:
    attrs = {**ATTRS, "unit_of_measurement": "Wh"}
    hass.states.async_set(ENTITY, "1000000", attrs)
    entry = await setup_entry(hass)

    await advance(hass, freezer, 5)
    hass.states.async_set(ENTITY, STATE_UNAVAILABLE)
    await hass.async_block_till_done()
    await advance(hass, freezer, 120)
    hass.states.async_set(ENTITY, "1012000", attrs)
    await hass.async_block_till_done()
    await advance(hass, freezer, 16)

    [session] = entry.runtime_data.sessions
    assert session.kwh == 12.0
    assert session.source == SOURCE_UNOBSERVED


async def test_sessions_survive_reload(
    recorder_mock, hass: HomeAssistant, freezer: FrozenDateTimeFactory
) -> None:
    hass.states.async_set(ENTITY, "1000.0", ATTRS)
    entry = await setup_entry(hass)
    await advance(hass, freezer, 1)
    hass.states.async_set(ENTITY, "1004.0", ATTRS)
    await hass.async_block_till_done()
    await advance(hass, freezer, 16)
    assert len(entry.runtime_data.sessions) == 1
    tracking_since = entry.runtime_data.tracking_since

    assert await hass.config_entries.async_reload(entry.entry_id)
    await hass.async_block_till_done()
    assert len(entry.runtime_data.sessions) == 1
    assert entry.runtime_data.tracking_since == tracking_since
    assert entry.runtime_data.tracker.last_reading == 1004.0


async def test_generate_report(
    recorder_mock,
    hass: HomeAssistant,
    hass_client: ClientSessionGenerator,
) -> None:
    await hass.config.async_set_time_zone("Europe/Amsterdam")
    hass.states.async_set(ENTITY, "1021.0", ATTRS)
    entry = await setup_entry(hass)

    year, quarter = previous_quarter(dt_util.now().date())
    start, end = quarter_bounds(year, quarter, dt_util.get_default_time_zone())
    first_hour = dt_util.as_utc(start) - timedelta(hours=1)
    changes = [0.0] * int((dt_util.as_utc(end) - first_hour) / timedelta(hours=1))
    changes[10], changes[11], changes[800] = 6.0, 4.0, 11.0
    await import_hourly(hass, first_hour, changes, 1000.0)

    events = []
    hass.bus.async_listen(EVENT_REPORT_GENERATED, events.append)
    response = await hass.services.async_call(
        DOMAIN,
        "generate_report",
        {"year": year, "quarter": quarter},
        blocking=True,
        return_response=True,
    )
    await hass.async_block_till_done()

    assert response["config_entry_id"] == entry.entry_id
    assert response["total_kwh"] == 21.0
    assert response["meter_begin"] == 1000.0
    assert response["meter_end"] == 1021.0
    assert response["sessions"] == 2
    assert response["complete"] is True
    assert len(events) == 1

    xlsx = Path(response["xlsx_path"])
    assert xlsx.name == f"ere_laadpaal_{year}_q{quarter}.xlsx"
    workbook = load_workbook(xlsx)
    sessions = list(workbook["Sessies"].iter_rows(values_only=True))
    assert sessions[1][2] == start.replace(tzinfo=None) + timedelta(hours=9)
    assert sessions[1][7] == 10.0
    assert Path(response["csv_path"]).is_file()

    client = await hass_client()
    download = await client.get(f"/api/{DOMAIN}/{xlsx.name}")
    assert download.status == 200
    assert (await client.get(f"/api/{DOMAIN}/secrets.yaml")).status == 404
    assert (await client.get(f"/api/{DOMAIN}/ere_x_2020_q1.csv")).status == 404

    # Without arguments the previous quarter is reported.
    response = await hass.services.async_call(
        DOMAIN, "generate_report", {}, blocking=True, return_response=True
    )
    assert (response["year"], response["quarter"]) == (year, quarter)

    xlsx.unlink()
    Path(response["csv_path"]).unlink()


async def test_automatic_report_on_quarter_change(
    recorder_mock, hass: HomeAssistant, freezer: FrozenDateTimeFactory
) -> None:
    # Timers only fire for moments after the real clock, so use future dates.
    await hass.config.async_set_time_zone("Europe/Amsterdam")
    freezer.move_to("2035-09-30 12:00:00+00:00")
    hass.states.async_set(ENTITY, "1000.0", ATTRS)
    entry = await setup_entry(hass)
    events = []
    hass.bus.async_listen(EVENT_REPORT_GENERATED, events.append)

    async def move_to(moment: str) -> None:
        freezer.move_to(moment)
        async_fire_time_changed(hass)
        await hass.async_block_till_done(wait_background_tasks=True)

    # The first check after installation only notes which quarter was last.
    await move_to("2035-09-30 22:31:00+00:00")
    assert events == []
    assert entry.runtime_data._last_auto_report == "2035Q3"

    await move_to("2035-10-01 22:31:00+00:00")
    assert events == []

    await move_to("2035-12-31 23:31:00+00:00")
    assert len(events) == 1
    assert (events[0].data["year"], events[0].data["quarter"]) == (2035, 4)
    assert entry.runtime_data._last_auto_report == "2035Q4"

    for name in ("xlsx_path", "csv_path"):
        Path(events[0].data[name]).unlink()
