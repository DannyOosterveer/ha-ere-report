"""Tests for reconstruction and report building."""

import csv
from datetime import UTC, datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from custom_components.ere_report.history import (
    HourRow,
    previous_quarter,
    quarter_bounds,
    reconstruct_sessions,
)
from custom_components.ere_report.report import (
    ReportMeta,
    build_report,
    write_csv,
    write_xlsx,
)
from custom_components.ere_report.session_tracker import (
    FLAG_SPLIT,
    SOURCE_LIVE,
    SOURCE_RECONSTRUCTED,
    Session,
)

TZ = ZoneInfo("Europe/Amsterdam")
META = ReportMeta(
    charger_name="Laadpaal",
    energy_entity="sensor.meter",
    holder_name="J. Jansen",
    address="Dorpsstraat 1",
    postcode_city="1234 AB Ons Dorp",
    ean="871234567890123456",
    charger_brand="Alfen",
    charger_model="Eve Single Pro-line",
    charger_serial="ACE0001",
    mid_confirmed=True,
)
Q3_START, Q3_END = quarter_bounds(2026, 3, TZ)


def hourly(start: datetime, changes: list[float], meter: float) -> list[HourRow]:
    """Build consecutive hour rows with a running meter state."""
    rows = []
    for index, change in enumerate(changes):
        meter = round(meter + change, 3)
        rows.append(HourRow(start + timedelta(hours=index), change, meter))
    return rows


def quarter_rows(changes_by_hour: dict[int, float], meter: float = 1000.0):
    """Hour rows for all of Q3 2026 plus one hour before it."""
    first = Q3_START.astimezone(UTC) - timedelta(hours=1)
    hours = int((Q3_END.astimezone(UTC) - first) / timedelta(hours=1))
    return hourly(first, [changes_by_hour.get(i, 0.0) for i in range(hours)], meter)


def test_quarter_helpers() -> None:
    start, end = quarter_bounds(2026, 4, TZ)
    assert (start.month, end.year, end.month) == (10, 2027, 1)
    assert previous_quarter(datetime(2026, 10, 1).date()) == (2026, 3)
    assert previous_quarter(datetime(2026, 2, 1).date()) == (2025, 4)


def test_reconstruct_clusters_hours() -> None:
    start = datetime(2026, 7, 4, 10, tzinfo=UTC)
    rows = hourly(start, [0, 5.0, 7.0, 0.01, 3.0, 0, 0, 2.0], meter=100.0)
    first, second = reconstruct_sessions(rows, 0.05)
    assert first.start == start + timedelta(hours=1)
    assert first.end == start + timedelta(hours=5)
    assert first.kwh == 15.0
    assert first.meter_start == 100.0
    assert first.meter_end == 115.01
    assert first.source == SOURCE_RECONSTRUCTED
    assert second.kwh == 2.0


def test_report_totals_and_reconstruction() -> None:
    rows = quarter_rows({10: 6.0, 11: 4.0, 800: 11.0, 1500: 0.02})
    now = Q3_END.astimezone(UTC) + timedelta(hours=1)
    data = build_report(META, 2026, 3, TZ, now, rows, [], None, 0.05)
    assert data.complete
    assert data.meter_begin == 1000.0
    assert data.meter_end == 1021.02
    assert data.meter_total == 21.02
    assert [s.kwh for s in data.sessions] == [10.0, 11.0]
    assert data.sessions_total == 21.0
    assert data.unallocated == 0.02
    assert any("gereconstrueerd" in note for note in data.notes)
    assert not any("geen meetgegevens" in note for note in data.notes)


def test_live_sessions_replace_reconstruction_after_install() -> None:
    rows = quarter_rows({10: 6.0, 800: 11.0})
    tracking_since = Q3_START.astimezone(UTC) + timedelta(hours=400)
    live_start = Q3_START.astimezone(UTC) + timedelta(hours=799, minutes=12)
    live = Session(
        start=live_start,
        end=live_start + timedelta(minutes=47),
        kwh=11.0,
        meter_start=1006.0,
        meter_end=1017.0,
    )
    now = Q3_END.astimezone(UTC) + timedelta(hours=1)
    data = build_report(META, 2026, 3, TZ, now, rows, [live], tracking_since, 0.05)
    assert [s.source for s in data.sessions] == [SOURCE_RECONSTRUCTED, SOURCE_LIVE]
    assert data.sessions_total == 17.0
    assert data.unallocated == 0.0


def test_session_across_quarter_start_is_split() -> None:
    rows = quarter_rows({0: 4.0, 1: 6.0})
    boundary = Q3_START.astimezone(UTC)
    live = Session(
        start=boundary - timedelta(minutes=40),
        end=boundary + timedelta(minutes=50),
        kwh=10.0,
        meter_start=1000.0,
        meter_end=1010.0,
    )
    now = Q3_END.astimezone(UTC) + timedelta(hours=1)
    data = build_report(
        META, 2026, 3, TZ, now, rows, [live], boundary - timedelta(days=30), 0.05
    )
    [session] = data.sessions
    assert session.start == Q3_START
    assert session.meter_start == 1004.0
    assert session.kwh == 6.0
    assert session.flags == [FLAG_SPLIT]
    assert data.unallocated == 0.0


def test_notes_for_missing_hours_and_spikes() -> None:
    rows = [r for i, r in enumerate(quarter_rows({20: 30.0})) if not 5 <= i < 8]
    now = Q3_END.astimezone(UTC) + timedelta(hours=1)
    meta = ReportMeta(charger_name="Laadpaal", energy_entity="sensor.meter")
    data = build_report(meta, 2026, 3, TZ, now, rows, [], None, 0.05)
    text = " ".join(data.notes)
    assert "3 van de 2208 uren" in text
    assert "30.0 kWh" in text
    assert "MID-meter" in text


def test_incomplete_quarter_is_marked() -> None:
    rows = quarter_rows({10: 6.0})[:200]
    now = rows[-1].start + timedelta(hours=1)
    data = build_report(META, 2026, 3, TZ, now, rows, [], None, 0.05)
    assert not data.complete
    assert data.notes[0].startswith("Voorlopig rapport")
    assert not any("geen meetgegevens" in note for note in data.notes)


def test_write_files(tmp_path: Path) -> None:
    rows = quarter_rows({10: 6.0, 11: 4.0, 800: 11.0})
    now = Q3_END.astimezone(UTC) + timedelta(hours=1)
    data = build_report(META, 2026, 3, TZ, now, rows, [], None, 0.05)
    xlsx = tmp_path / "out" / "report.xlsx"
    write_xlsx(data, xlsx, TZ, "0.1.0")
    write_csv(data, xlsx.with_suffix(".csv"), TZ)

    workbook = load_workbook(xlsx)
    assert workbook.sheetnames == ["Samenvatting", "Sessies", "Maandtotalen"]
    summary = {
        row[0]: row[1] for row in workbook["Samenvatting"].iter_rows(values_only=True)
    }
    assert summary["EAN-code aansluiting"] == "871234567890123456"
    assert summary["Meterstand begin periode (kWh)"] == 1000.0
    assert summary["Meterstand eind periode (kWh)"] == 1021.0
    assert summary["Geleverd in periode (kWh)"] == 21.0
    assert summary["Aantal laadsessies"] == 2
    sessions = list(workbook["Sessies"].iter_rows(values_only=True))
    assert sessions[1][2] == datetime(2026, 7, 1, 9, 0)
    assert sessions[1][7] == 10.0
    months = list(workbook["Maandtotalen"].iter_rows(values_only=True))
    assert months[1] == ("juli", 1, 10.0, 10.0)
    assert months[2] == ("augustus", 1, 11.0, 11.0)

    with xlsx.with_suffix(".csv").open(encoding="utf-8") as handle:
        lines = list(csv.DictReader(handle))
    assert len(lines) == 2
    assert lines[0]["start"] == "2026-07-01T09:00:00+02:00"
    assert lines[0]["kwh"] == "10.00"
    assert lines[0]["ean"] == "871234567890123456"
